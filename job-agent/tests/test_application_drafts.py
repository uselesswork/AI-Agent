from io import BytesIO
from pathlib import Path

import pymupdf
from docx import Document
from fastapi.testclient import TestClient

from app.main import app
from app.models.draft import DraftDocument, DraftSection
from app.services.draft_exporter import export_pdf
from tests.test_applications_api import application_payload, use_temporary_repository

client = TestClient(app)


def test_draft_uses_saved_facts_and_separates_suggestions(monkeypatch, tmp_path: Path) -> None:
    use_temporary_repository(monkeypatch, tmp_path)
    payload = application_payload()
    payload["candidate_profile"]["skills"] = [
        {"name": "SQL", "category": "数据库"},
        {"name": "Python", "category": "编程语言"},
    ]
    payload["candidate_profile"]["projects"] = [
        {"name": "课程管理系统", "description": "使用 Python 开发课程管理接口。"},
    ]
    payload["job_profile"]["required_skills"] = ["Python"]
    payload["match_result"]["resume_suggestions"] = ["补充量化成果"]
    record_id = client.post("/api/applications", json=payload).json()["id"]

    response = client.get(f"/api/applications/{record_id}/draft")
    assert response.status_code == 200
    result = response.json()
    skills = next(section["lines"] for section in result["resume"]["sections"]
                  if section["heading"] == "技能")
    assert skills == ["Python", "SQL"]
    assert result["suggestions"] == ["补充量化成果"]
    document_text = str(result["resume"]) + str(result["cover_letter"])
    assert "补充量化成果" not in document_text
    assert "课程管理系统" in document_text
    assert "投递前请逐项核对原件" in result["warnings"][0]


def test_draft_exports_docx_and_pdf(monkeypatch, tmp_path: Path) -> None:
    use_temporary_repository(monkeypatch, tmp_path)
    record_id = client.post("/api/applications", json=application_payload()).json()["id"]
    for kind in ("resume", "letter"):
        docx_response = client.get(f"/api/applications/{record_id}/draft/{kind}.docx")
        assert docx_response.status_code == 200
        assert docx_response.content.startswith(b"PK")
        assert "attachment" in docx_response.headers["content-disposition"]
        document = Document(BytesIO(docx_response.content))
        assert any("张晨" in paragraph.text for paragraph in document.paragraphs)

        pdf_response = client.get(f"/api/applications/{record_id}/draft/{kind}.pdf")
        assert pdf_response.status_code == 200
        assert pdf_response.content.startswith(b"%PDF")
        with pymupdf.open(stream=pdf_response.content, filetype="pdf") as pdf:
            assert "张晨" in "".join(page.get_text() for page in pdf)


def test_missing_application_draft_returns_404(monkeypatch, tmp_path: Path) -> None:
    use_temporary_repository(monkeypatch, tmp_path)
    assert client.get("/api/applications/999/draft").status_code == 404
    assert client.get("/api/applications/999/draft/resume.pdf").status_code == 404


def test_pdf_wraps_long_content_across_pages() -> None:
    draft = DraftDocument(title="长内容排版测试", sections=[
        DraftSection(heading="项目经历", lines=["中文内容需要自动换行。" * 600]),
    ])
    with pymupdf.open(stream=export_pdf(draft), filetype="pdf") as pdf:
        assert len(pdf) > 1
        assert all(page.get_text().strip() for page in pdf)
