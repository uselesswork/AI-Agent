import os
from io import BytesIO
from pathlib import Path

import pymupdf
from docx import Document
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from docx.shared import RGBColor

from app.models.draft import DraftDocument


class PdfFontUnavailableError(RuntimeError):
    """系统没有可用的中文 TrueType 字体。"""


def export_docx(draft: DraftDocument) -> bytes:
    document = Document()
    section = document.sections[0]
    section.top_margin = section.bottom_margin = Cm(2)
    section.left_margin = section.right_margin = Cm(2.2)
    normal = document.styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal.font.size = Pt(10.5)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    for style_name in ("Title", "Heading 1"):
        style = document.styles[style_name]
        style.font.name = "Microsoft YaHei"
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    document.add_heading(draft.title, level=0)
    for part in draft.sections:
        document.add_heading(part.heading, level=1)
        for line in part.lines:
            document.add_paragraph(line)
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _cjk_font_path() -> Path:
    configured = os.getenv("PDF_CJK_FONT_PATH")
    candidates = ([Path(configured)] if configured else []) + [
        Path("C:/Windows/Fonts/NotoSansSC-VF.ttf"),
        Path("C:/Windows/Fonts/simhei.ttf"),
        Path("/usr/share/fonts/truetype/noto/NotoSansCJK-Regular.ttc"),
        Path("/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc"),
    ]
    for path in candidates:
        if path.is_file():
            return path
    raise PdfFontUnavailableError("未找到中文字体；请设置 PDF_CJK_FONT_PATH 指向中文 TTF/TTC 字体。")


def _wrap(line: str, font: pymupdf.Font, size: float, max_width: float) -> list[str]:
    if not line:
        return [""]
    result: list[str] = []
    current = ""
    for character in line:
        if current and font.text_length(current + character, fontsize=size) > max_width:
            result.append(current)
            current = character
        else:
            current += character
    if current:
        result.append(current)
    return result


def export_pdf(draft: DraftDocument) -> bytes:
    font_path = _cjk_font_path()
    font = pymupdf.Font(fontfile=str(font_path))
    document = pymupdf.open()
    page = document.new_page(width=595, height=842)
    page.insert_font(fontname="draft_cjk", fontfile=str(font_path))
    y = 65.0

    def write(line: str, size: float, step: float) -> None:
        nonlocal page, y
        for wrapped in _wrap(line, font, size, 495):
            if y + step > 790:
                page = document.new_page(width=595, height=842)
                page.insert_font(fontname="draft_cjk", fontfile=str(font_path))
                y = 65.0
            if wrapped:
                page.insert_text((50, y), wrapped, fontname="draft_cjk", fontsize=size,
                                 color=(0.15, 0.2, 0.17))
            y += step

    write(draft.title, 17, 27)
    y += 8
    for part in draft.sections:
        write(part.heading, 12, 20)
        for line in part.lines:
            write(line, 10.5, 17)
        y += 9
    result = document.tobytes(garbage=4, deflate=True)
    document.close()
    return result
